"""Write data/jobs.xlsx with 'New Jobs' and 'All Jobs' sheets."""
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

NEW_COLS = ["Score", "Title", "Company", "Location", "Posted Date", "Source", "URL", "First Seen"]
ALL_COLS = NEW_COLS + ["Last Seen"]
WIDTHS = {"Score": 7, "Title": 45, "Company": 22, "Location": 25, "Posted Date": 12,
          "Source": 16, "URL": 60, "First Seen": 20, "Last Seen": 20}


def _frame(rows: list[dict], cols: list[str]) -> pd.DataFrame:
    df = pd.DataFrame([{
        "Score": r["score"], "Title": r["title"], "Company": r["company"],
        "Location": r["location"], "Posted Date": r["posted_date"], "Source": r["source"],
        "URL": r["url"], "First Seen": r["first_seen"], "Last Seen": r["last_seen"],
    } for r in rows], columns=ALL_COLS)
    return df[cols].sort_values("Score", ascending=False, kind="stable")


def write_workbook(path: str, new_rows: list[dict], all_rows: list[dict]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, rows, cols in [("New Jobs", new_rows, NEW_COLS), ("All Jobs", all_rows, ALL_COLS)]:
            df = _frame(rows, cols)
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(cols, 1):
                ws.column_dimensions[get_column_letter(i)].width = WIDTHS[col]
            url_col = cols.index("URL") + 1
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=url_col)
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
